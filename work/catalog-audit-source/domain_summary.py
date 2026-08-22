from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


SOURCE = Path(r"C:\Users\Admin\Downloads\New folder (3)")
OUT = Path(r"C:\alumdoor\work\catalog-audit-source\domain_summary.json")


def norm(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip().casefold()


def display(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if hasattr(v, "text"):
        return getattr(v, "text")
    return v


def counter_dict(values, limit=None):
    result = Counter(norm(v) for v in values if norm(v))
    rows = [{"value": k, "count": n} for k, n in result.most_common(limit)]
    return rows


def dup_map(records: list[tuple[int, Any]]) -> list[dict[str, Any]]:
    seen: dict[str, list[tuple[int, Any]]] = defaultdict(list)
    for row, value in records:
        key = norm(value)
        if key:
            seen[key].append((row, value))
    out = []
    for key, rows in seen.items():
        if len(rows) > 1:
            out.append({"normalized": key, "rows": [r for r, _ in rows], "values": [display(v) for _, v in rows]})
    return sorted(out, key=lambda x: (-len(x["rows"]), x["normalized"]))


def parse_money(v: Any) -> int | None:
    if isinstance(v, (int, float)):
        return round(v)
    if not isinstance(v, str):
        return None
    digits = re.sub(r"[^0-9-]", "", v)
    if not digits or digits == "-":
        return None
    return int(digits)


def product_catalog():
    path = SOURCE / "danh mục sản phẩm.xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Sheet1"]
    headers = [ws.cell(1, c).value for c in range(1, 12)]
    data_end_row = max(
        r for r in range(1, ws.max_row + 1)
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, 12))
    )
    rows = []
    for r in range(2, data_end_row + 1):
        row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, 12)}
        row["_row"] = r
        rows.append(row)
    required = ["Mã SP", "TÊN SP", "Nhóm SP", "ĐVT", "Giá niêm yết"]
    blanks = {h: [row["_row"] for row in rows if row.get(h) in (None, "")] for h in headers}
    parsed = []
    invalid_prices = []
    deltas = []
    for row in rows:
        normal = parse_money(row.get("Giá niêm yết"))
        rail = parse_money(row.get("Giá có ray"))
        if row.get("Giá niêm yết") not in (None, "") and normal is None:
            invalid_prices.append({"row": row["_row"], "field": "Giá niêm yết", "value": row.get("Giá niêm yết")})
        if row.get("Giá có ray") not in (None, "") and rail is None:
            invalid_prices.append({"row": row["_row"], "field": "Giá có ray", "value": row.get("Giá có ray")})
        if normal is not None:
            parsed.append(normal)
        if normal is not None and rail is not None and row.get("Có ray tặng") is True:
            deltas.append({"row": row["_row"], "code": row.get("Mã SP"), "delta": rail - normal})
    delta_counts = Counter(d["delta"] for d in deltas)
    outliers = [d for d in deltas if d["delta"] not in (0, 75000)]
    return {
        "headers": headers,
        "record_count": len(rows),
        "required_blank_rows": {k: v for k, v in blanks.items() if k in required and v},
        "all_blank_counts": {k: len(v) for k, v in blanks.items()},
        "duplicate_codes": dup_map([(row["_row"], row.get("Mã SP")) for row in rows]),
        "duplicate_names": dup_map([(row["_row"], row.get("TÊN SP")) for row in rows]),
        "groups": counter_dict([r.get("Nhóm SP") for r in rows]),
        "uoms": counter_dict([r.get("ĐVT") for r in rows]),
        "logic_groups": counter_dict([r.get("Nhóm logic") for r in rows]),
        "rail_flags": counter_dict([r.get("Có ray tặng") for r in rows]),
        "ultra_small_price_nonblank": [
            {"row": r["_row"], "code": r.get("Mã SP"), "name": r.get("TÊN SP"), "value": r.get("PT siêu nhỏ (đ/bộ)")}
            for r in rows if r.get("PT siêu nhỏ (đ/bộ)") not in (None, "")
        ],
        "price_min": min(parsed) if parsed else None,
        "price_max": max(parsed) if parsed else None,
        "nonpositive_prices": [
            {"row": r["_row"], "code": r.get("Mã SP"), "name": r.get("TÊN SP"), "value": r.get("Giá niêm yết")}
            for r in rows if (parse_money(r.get("Giá niêm yết")) or 0) <= 0 and r.get("Giá niêm yết") not in (None, "")
        ],
        "invalid_prices": invalid_prices,
        "rail_price_delta_counts": [{"delta": k, "count": v} for k, v in delta_counts.most_common()],
        "rail_price_delta_outliers": outliers,
        "rows": [{k: display(v) for k, v in row.items()} for row in rows],
    }


def linked_workbook():
    path = SOURCE / "MS LIÊN BS.xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)

    dm = wb["DANH MỤC"]
    customer_rows = []
    for r in range(4, dm.max_row + 1):
        if any(dm.cell(r, c).value not in (None, "") for c in range(1, 5)):
            customer_rows.append({"row": r, "tax_id": dm.cell(r, 1).value, "name": dm.cell(r, 2).value, "owner": dm.cell(r, 3).value, "type": dm.cell(r, 4).value})
    product_rows = []
    for r in range(5, dm.max_row + 1):
        if any(dm.cell(r, c).value not in (None, "") for c in range(6, 16)):
            product_rows.append({
                "row": r,
                "price_date": dm.cell(r, 6).value,
                "sales_code": dm.cell(r, 7).value,
                "purchase_code": dm.cell(r, 8).value,
                "purchase_price": dm.cell(r, 9).value,
                "purchase_uom": dm.cell(r, 10).value,
                "sales_price": dm.cell(r, 11).value,
                "sales_uom": dm.cell(r, 12).value,
                "supplier": dm.cell(r, 13).value,
                "note": dm.cell(r, 14).value,
                "cost_formula": display(dm.cell(r, 15).value),
            })

    item_ws = wb["ĐM"]
    item_rows = []
    for r in range(7, item_ws.max_row + 1):
        name = item_ws.cell(r, 3).value
        code = item_ws.cell(r, 4).value
        if name not in (None, "") or code not in (None, ""):
            item_rows.append({
                "row": r,
                "group": item_ws.cell(r, 1).value,
                "ordinal": item_ws.cell(r, 2).value,
                "name": display(name),
                "code": display(code),
                "finished_product": display(item_ws.cell(r, 5).value),
                "uom": display(item_ws.cell(r, 6).value),
                "norm": display(item_ws.cell(r, 7).value),
                "sales_price": display(item_ws.cell(r, 8).value),
                "formula_type": display(item_ws.cell(r, 23).value),
            })

    stock_ws = wb["Trang tính29"]
    stock_rows = []
    for r in range(2, stock_ws.max_row + 1):
        if any(stock_ws.cell(r, c).value not in (None, "") for c in range(1, 8)):
            stock_rows.append({
                "row": r,
                "type": stock_ws.cell(r, 1).value,
                "code": stock_ws.cell(r, 2).value,
                "name": stock_ws.cell(r, 3).value,
                "color": stock_ws.cell(r, 4).value,
                "qty_formula": display(stock_ws.cell(r, 5).value),
                "qty_cached": wbv["Trang tính29"].cell(r, 5).value,
                "uom": stock_ws.cell(r, 6).value,
                "check_date": display(stock_ws.cell(r, 7).value),
            })

    tx_ws = wb["chi tiết nhập hàng ngày"]
    transaction_rows = []
    for r in range(6, tx_ws.max_row + 1):
        if any(tx_ws.cell(r, c).value not in (None, "") for c in range(1, 30)):
            transaction_rows.append({
                "row": r,
                "date_parts": [display(tx_ws.cell(r, c).value) for c in (1, 2, 3)],
                "document": tx_ws.cell(r, 5).value,
                "party": tx_ws.cell(r, 6).value,
                "transaction_type": tx_ws.cell(r, 8).value,
                "item_name": tx_ws.cell(r, 9).value,
                "item_code": display(tx_ws.cell(r, 10).value),
                "uom": display(tx_ws.cell(r, 11).value),
                "status": tx_ws.cell(r, 29).value,
            })

    return {
        "customers": {
            "count": len(customer_rows),
            "blank_tax_id_count": sum(not norm(r["tax_id"]) for r in customer_rows),
            "blank_name_rows": [r["row"] for r in customer_rows if not norm(r["name"])],
            "duplicate_names": dup_map([(r["row"], r["name"]) for r in customer_rows]),
            "types": counter_dict([r["type"] for r in customer_rows]),
            "owners": counter_dict([r["owner"] for r in customer_rows]),
        },
        "price_crosswalk": {
            "count": len(product_rows),
            "blank_sales_code_rows": [r["row"] for r in product_rows if not norm(r["sales_code"])],
            "blank_purchase_code_rows": [r["row"] for r in product_rows if not norm(r["purchase_code"])],
            "duplicate_sales_codes": dup_map([(r["row"], r["sales_code"]) for r in product_rows]),
            "duplicate_purchase_codes": dup_map([(r["row"], r["purchase_code"]) for r in product_rows]),
            "sales_uoms": counter_dict([r["sales_uom"] for r in product_rows]),
            "purchase_uoms": counter_dict([r["purchase_uom"] for r in product_rows]),
            "supplier_count": len({norm(r["supplier"]) for r in product_rows if norm(r["supplier"])}),
            "cost_formula_count": sum(r["cost_formula"] not in (None, "") for r in product_rows),
            "rows": [{k: display(v) for k, v in r.items()} for r in product_rows],
        },
        "item_master": {
            "count": len(item_rows),
            "duplicate_names": dup_map([(r["row"], r["name"]) for r in item_rows]),
            "duplicate_codes": dup_map([(r["row"], r["code"]) for r in item_rows]),
            "blank_name_rows": [r["row"] for r in item_rows if not norm(r["name"])],
            "blank_code_rows": [r["row"] for r in item_rows if not norm(r["code"])],
            "groups": counter_dict([r["group"] for r in item_rows]),
            "uoms": counter_dict([r["uom"] for r in item_rows]),
            "formula_types": counter_dict([r["formula_type"] for r in item_rows]),
            "rows": item_rows,
        },
        "stock_snapshot": {
            "count": len(stock_rows),
            "duplicate_codes": dup_map([(r["row"], r["code"]) for r in stock_rows]),
            "blank_codes": [r["row"] for r in stock_rows if not norm(r["code"])],
            "types": counter_dict([r["type"] for r in stock_rows]),
            "colors": counter_dict([r["color"] for r in stock_rows]),
            "uoms": counter_dict([r["uom"] for r in stock_rows]),
            "formula_qty_count": sum(isinstance(r["qty_formula"], str) and r["qty_formula"].startswith("=") for r in stock_rows),
            "rows": stock_rows,
        },
        "transactions": {
            "count": len(transaction_rows),
            "transaction_types": counter_dict([r["transaction_type"] for r in transaction_rows]),
            "statuses": counter_dict([r["status"] for r in transaction_rows]),
            "blank_document_count": sum(not norm(r["document"]) for r in transaction_rows),
            "blank_party_count": sum(not norm(r["party"]) for r in transaction_rows),
            "blank_item_name_count": sum(not norm(r["item_name"]) for r in transaction_rows),
        },
    }


def inventory_workbook():
    path = SOURCE / "TỒN NHÔM 2026 NEW (1).xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    operational = []
    for ws in wb.worksheets:
        if ws.title in ("LICH_SU", "LỊCH SỬ"):
            continue
        records = []
        for r in range(10, ws.max_row + 1):
            length = ws.cell(r, 4).value
            leaves = ws.cell(r, 5).value
            if any(ws.cell(r, c).value not in (None, "") for c in range(1, 7)):
                records.append({
                    "row": r,
                    "date": display(ws.cell(r, 1).value),
                    "color_or_type": ws.cell(r, 2).value,
                    "condition_or_color": ws.cell(r, 3).value,
                    "length": length,
                    "leaves": leaves,
                    "return_date": display(ws.cell(r, 6).value),
                    "status_cached": wbv[ws.title].cell(r, 7).value,
                    "selected": ws.cell(r, 8).value,
                    "scrap_cached": wbv[ws.title].cell(r, 9).value,
                    "kg_total": ws.cell(r, 10).value,
                    "entry_note": ws.cell(r, 11).value,
                    "note": ws.cell(r, 12).value,
                })
        business_keys = []
        for rec in records:
            key = "|".join(norm(rec[k]) for k in ("date", "color_or_type", "condition_or_color", "length", "leaves", "return_date"))
            business_keys.append((rec["row"], key))
        numeric = [r for r in records if isinstance(r["length"], (int, float)) and isinstance(r["leaves"], (int, float))]
        operational.append({
            "sheet": ws.title,
            "records": len(records),
            "numeric_stock_rows": len(numeric),
            "total_leaves": sum(r["leaves"] for r in numeric),
            "total_linear_m_leaf": sum(r["length"] * r["leaves"] for r in numeric),
            "zero_or_negative_length_rows": [r["row"] for r in numeric if r["length"] <= 0],
            "zero_or_negative_leaves_rows": [r["row"] for r in numeric if r["leaves"] <= 0],
            "scrap_length_under_015_rows": [r["row"] for r in numeric if r["length"] < 0.15],
            "selected_rows": [r["row"] for r in records if r["selected"] is True],
            "colors_or_types": counter_dict([r["color_or_type"] for r in records]),
            "conditions_or_colors": counter_dict([r["condition_or_color"] for r in records]),
            "status_cached": counter_dict([r["status_cached"] for r in records]),
            "duplicate_business_keys": dup_map(business_keys),
        })

    log = wb["LICH_SU"]
    logs = []
    for r in range(2, log.max_row + 1):
        if any(log.cell(r, c).value not in (None, "") for c in range(1, 8)):
            logs.append({
                "row": r,
                "time": display(log.cell(r, 1).value),
                "customer": log.cell(r, 2).value,
                "document": log.cell(r, 3).value,
                "action": log.cell(r, 4).value,
                "detail": log.cell(r, 5).value,
                "source_row": log.cell(r, 6).value,
                "item_code": log.cell(r, 7).value,
            })

    history = wb["LỊCH SỬ"]
    histories = []
    for r in range(2, history.max_row + 1):
        if any(history.cell(r, c).value not in (None, "") for c in range(1, 22)):
            histories.append({
                "row": r,
                "direction": history.cell(r, 1).value,
                "date": display(history.cell(r, 2).value),
                "type": history.cell(r, 3).value,
                "color": history.cell(r, 4).value,
                "condition": history.cell(r, 5).value,
                "length": history.cell(r, 6).value,
                "leaves": history.cell(r, 7).value,
                "cut_length": history.cell(r, 8).value,
                "cut_leaves": history.cell(r, 9).value,
                "remainder_length": history.cell(r, 10).value,
                "remainder_leaves": history.cell(r, 11).value,
                "customer": history.cell(r, 12).value,
                "source_receipt": history.cell(r, 13).value,
                "document": history.cell(r, 14).value,
                "returned": history.cell(r, 15).value,
                "reason": history.cell(r, 16).value,
                "operation_code": history.cell(r, 17).value,
                "source_sheet": history.cell(r, 18).value,
                "error_note": history.cell(r, 19).value,
                "source_detail": history.cell(r, 20).value,
                "remainder_row": history.cell(r, 21).value,
            })
    return {
        "operational_sheets": operational,
        "event_log": {
            "count": len(logs),
            "actions": counter_dict([r["action"] for r in logs]),
            "blank_customer": sum(not norm(r["customer"]) for r in logs),
            "blank_document": sum(not norm(r["document"]) for r in logs),
            "blank_item_code": sum(not norm(r["item_code"]) for r in logs),
            "duplicate_exact_events": dup_map([(r["row"], "|".join(norm(r[k]) for k in ("time", "customer", "document", "action", "detail", "source_row", "item_code"))) for r in logs]),
            "zero_cut_examples": [r for r in logs if re.search(r"cắt\s+[^x]*x0\b", norm(r["detail"]))][:20],
        },
        "history": {
            "count": len(histories),
            "directions": counter_dict([r["direction"] for r in histories]),
            "types": counter_dict([r["type"] for r in histories]),
            "colors": counter_dict([r["color"] for r in histories]),
            "conditions": counter_dict([r["condition"] for r in histories]),
            "blank_document": sum(not norm(r["document"]) for r in histories),
            "blank_operation_code": sum(not norm(r["operation_code"]) for r in histories),
            "duplicate_operation_codes": dup_map([(r["row"], r["operation_code"]) for r in histories]),
            "source_sheets": counter_dict([r["source_sheet"] for r in histories]),
            "returned_flags": counter_dict([r["returned"] for r in histories]),
            "reason_count": sum(bool(norm(r["reason"])) for r in histories),
            "error_note_count": sum(bool(norm(r["error_note"])) for r in histories),
        },
    }


def cross_coverage(product, linked, inventory):
    catalog_codes = {norm(r["Mã SP"]) for r in product["rows"] if norm(r["Mã SP"])}
    catalog_names = {norm(r["TÊN SP"]) for r in product["rows"] if norm(r["TÊN SP"])}
    item_rows = linked["item_master"]["rows"]
    item_codes = {norm(r["code"]) for r in item_rows if norm(r["code"])}
    item_names = {norm(r["name"]) for r in item_rows if norm(r["name"])}
    snapshot_rows = linked["stock_snapshot"]["rows"]
    snapshot_codes = {norm(r["code"]) for r in snapshot_rows if norm(r["code"])}
    inv_sheet_codes = {norm(r["sheet"]) for r in inventory["operational_sheets"] if norm(r["sheet"])}
    return {
        "catalog_code_in_item_master": len(catalog_codes & item_codes),
        "catalog_name_in_item_master": len(catalog_names & item_names),
        "catalog_total": len(catalog_codes),
        "catalog_codes_missing_item_master_sample": sorted(catalog_codes - item_codes)[:100],
        "catalog_names_missing_item_master_sample": sorted(catalog_names - item_names)[:100],
        "stock_snapshot_codes_in_item_master": len(snapshot_codes & item_codes),
        "stock_snapshot_total_codes": len(snapshot_codes),
        "inventory_sheet_codes_in_catalog_codes": len(inv_sheet_codes & catalog_codes),
        "inventory_operational_sheet_count": len(inv_sheet_codes),
        "inventory_sheet_codes_missing_catalog": sorted(inv_sheet_codes - catalog_codes),
    }


def main():
    product = product_catalog()
    linked = linked_workbook()
    inventory = inventory_workbook()
    result = {
        "product_catalog": product,
        "linked_workbook": linked,
        "inventory_workbook": inventory,
        "cross_coverage": cross_coverage(product, linked, inventory),
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=display), encoding="utf-8")
    print(json.dumps({"output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
