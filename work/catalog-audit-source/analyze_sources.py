from __future__ import annotations

import json
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\Admin\Downloads\New folder (3)")
OUT = Path(r"C:\alumdoor\work\catalog-audit-source")


def clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)
        return value
    if isinstance(value, (str, int, bool)):
        return value
    return str(value)


def norm(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


def row_payload(ws, row_idx: int, min_col: int, max_col: int) -> dict[str, Any]:
    values = [clean_value(ws.cell(row_idx, c).value) for c in range(min_col, max_col + 1)]
    return {
        "row": row_idx,
        "nonempty": sum(v not in (None, "") for v in values),
        "values": values,
    }


def analyze_workbook(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    wb_values = openpyxl.load_workbook(path, data_only=True, read_only=False)
    result: dict[str, Any] = {
        "file": path.name,
        "bytes": path.stat().st_size,
        "sheet_order": wb.sheetnames,
        "defined_names": [],
        "sheets": [],
    }

    for dn in wb.defined_names.values():
        result["defined_names"].append({
            "name": dn.name,
            "attr_text": dn.attr_text,
            "hidden": dn.hidden,
        })

    for ws in wb.worksheets:
        value_ws = wb_values[ws.title]
        nonempty = []
        formulas = []
        error_cells = []
        whitespace_cells = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is not None and value != "":
                    nonempty.append((cell.row, cell.column))
                    if isinstance(value, str) and value.startswith("="):
                        formulas.append({
                            "cell": cell.coordinate,
                            "formula": value,
                            "cached": clean_value(value_ws[cell.coordinate].value),
                        })
                    if isinstance(value, str) and value.strip() == "":
                        whitespace_cells.append(cell.coordinate)
                    cached = value_ws[cell.coordinate].value
                    if isinstance(cached, str) and cached.startswith("#"):
                        error_cells.append({"cell": cell.coordinate, "value": cached})

        if nonempty:
            min_row = min(r for r, _ in nonempty)
            max_row = max(r for r, _ in nonempty)
            min_col = min(c for _, c in nonempty)
            max_col = max(c for _, c in nonempty)
        else:
            min_row = max_row = min_col = max_col = 0

        row_counts = []
        blank_rows = []
        exact_rows: dict[tuple[str, ...], list[int]] = defaultdict(list)
        if nonempty:
            for r in range(min_row, max_row + 1):
                vals = [ws.cell(r, c).value for c in range(min_col, max_col + 1)]
                count = sum(v not in (None, "") for v in vals)
                row_counts.append((r, count))
                if count == 0:
                    blank_rows.append(r)
                else:
                    exact_rows[tuple(norm(v) for v in vals)].append(r)

        duplicate_rows = [
            {"rows": rows, "normalized_values": list(key)}
            for key, rows in exact_rows.items()
            if len(rows) > 1
        ]
        duplicate_rows.sort(key=lambda item: (-len(item["rows"]), item["rows"][0]))

        column_profiles = []
        if nonempty:
            for c in range(min_col, max_col + 1):
                items = []
                blanks = 0
                formulas_in_col = 0
                for r in range(min_row, max_row + 1):
                    value = ws.cell(r, c).value
                    if value in (None, ""):
                        blanks += 1
                    else:
                        items.append(norm(value))
                        if isinstance(value, str) and value.startswith("="):
                            formulas_in_col += 1
                counts = Counter(items)
                duplicates = [
                    {"value": v, "count": n}
                    for v, n in counts.most_common(12)
                    if n > 1 and v != ""
                ]
                column_profiles.append({
                    "column": openpyxl.utils.get_column_letter(c),
                    "nonempty": len(items),
                    "blank_within_used_rows": blanks,
                    "unique_normalized": len(counts),
                    "formula_count": formulas_in_col,
                    "duplicate_values": duplicates,
                })

        candidate_header_rows = []
        for r, count in row_counts[: min(len(row_counts), 60)]:
            values = [ws.cell(r, c).value for c in range(min_col, max_col + 1)]
            text_count = sum(isinstance(v, str) and not v.startswith("=") for v in values if v not in (None, ""))
            distinct = len({norm(v) for v in values if v not in (None, "")})
            candidate_header_rows.append({
                "row": r,
                "nonempty": count,
                "text": text_count,
                "distinct": distinct,
                "values": [clean_value(v) for v in values],
            })
        candidate_header_rows.sort(key=lambda x: (-x["text"], -x["distinct"], x["row"]))

        first_rows = []
        last_rows = []
        dense_rows = []
        if nonempty:
            first_nonempty_rows = [r for r, count in row_counts if count > 0][:30]
            last_nonempty_rows = [r for r, count in row_counts if count > 0][-8:]
            densest = sorted(row_counts, key=lambda x: (-x[1], x[0]))[:10]
            first_rows = [row_payload(ws, r, min_col, max_col) for r in first_nonempty_rows]
            last_rows = [row_payload(ws, r, min_col, max_col) for r in last_nonempty_rows]
            dense_rows = [row_payload(ws, r, min_col, max_col) for r, _ in densest]

        result["sheets"].append({
            "name": ws.title,
            "state": ws.sheet_state,
            "declared_dimension": ws.calculate_dimension(),
            "actual_range": (
                f"{openpyxl.utils.get_column_letter(min_col)}{min_row}:"
                f"{openpyxl.utils.get_column_letter(max_col)}{max_row}"
                if nonempty else None
            ),
            "actual_nonempty_cells": len(nonempty),
            "actual_rows": (max_row - min_row + 1) if nonempty else 0,
            "actual_cols": (max_col - min_col + 1) if nonempty else 0,
            "min_row": min_row,
            "max_row": max_row,
            "min_col": min_col,
            "max_col": max_col,
            "blank_rows_within_used_range_count": len(blank_rows),
            "blank_rows_within_used_range_sample": blank_rows[:50],
            "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
            "hidden_rows": [r for r, dim in ws.row_dimensions.items() if dim.hidden],
            "hidden_columns": [c for c, dim in ws.column_dimensions.items() if dim.hidden],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
            "tables": [
                {"name": table.name, "ref": table.ref, "style": table.tableStyleInfo.name if table.tableStyleInfo else None}
                for table in ws.tables.values()
            ],
            "data_validations": [
                {
                    "sqref": str(dv.sqref),
                    "type": dv.type,
                    "formula1": dv.formula1,
                    "formula2": dv.formula2,
                }
                for dv in ws.data_validations.dataValidation
            ],
            "formula_count": len(formulas),
            "formula_sample": formulas[:30],
            "formula_errors": error_cells[:100],
            "whitespace_only_cells": whitespace_cells[:100],
            "duplicate_row_group_count": len(duplicate_rows),
            "duplicate_rows_sample": duplicate_rows[:30],
            "column_profiles": column_profiles,
            "candidate_headers": candidate_header_rows[:12],
            "first_rows": first_rows,
            "last_rows": last_rows,
            "densest_rows": dense_rows,
        })

    wb.close()
    wb_values.close()
    return result


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def analyze_docx(path: Path) -> dict[str, Any]:
    doc = Document(path)
    blocks = []
    for idx, block in enumerate(doc.iter_inner_content(), start=1):
        if isinstance(block, Paragraph):
            text = paragraph_text(block)
            if text:
                blocks.append({
                    "index": idx,
                    "type": "paragraph",
                    "style": block.style.name if block.style else None,
                    "text": text,
                })
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                rows.append(["\n".join(p.text for p in cell.paragraphs).strip() for cell in row.cells])
            blocks.append({
                "index": idx,
                "type": "table",
                "rows": len(rows),
                "cols": max((len(r) for r in rows), default=0),
                "data": rows,
            })

    headers = []
    footers = []
    for sec_idx, section in enumerate(doc.sections, start=1):
        headers.append({
            "section": sec_idx,
            "text": [p.text.strip() for p in section.header.paragraphs if p.text.strip()],
        })
        footers.append({
            "section": sec_idx,
            "text": [p.text.strip() for p in section.footer.paragraphs if p.text.strip()],
        })

    return {
        "file": path.name,
        "bytes": path.stat().st_size,
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "last_modified_by": doc.core_properties.last_modified_by,
            "created": clean_value(doc.core_properties.created),
            "modified": clean_value(doc.core_properties.modified),
        },
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "inline_shape_count": len(doc.inline_shapes),
        "headers": headers,
        "footers": footers,
        "blocks": blocks,
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    workbook_results = []
    for path in sorted(SOURCE.glob("*.xlsx"), key=lambda p: p.name.casefold()):
        workbook_results.append(analyze_workbook(path))
    docx_results = [analyze_docx(path) for path in sorted(SOURCE.glob("*.docx"))]
    (OUT / "xlsx_audit.json").write_text(
        json.dumps(workbook_results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "docx_audit.json").write_text(
        json.dumps(docx_results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({
        "xlsx_files": len(workbook_results),
        "docx_files": len(docx_results),
        "xlsx_output": str(OUT / "xlsx_audit.json"),
        "docx_output": str(OUT / "docx_audit.json"),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
